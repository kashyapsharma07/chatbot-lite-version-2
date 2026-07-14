#!/usr/bin/env python3
"""
API Costing Quotation v4.1 — All amounts in INR (₹)
Daily/Monthly, Two Scenarios, ₹ primary / $ secondary
"""

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml
import datetime

# ── DESIGN TOKENS ─────────────────────────────────────────────────────
NAVY=RGBColor(0x0F,0x17,0x2A); BLUE=RGBColor(0x2D,0x5B,0xE3)
TEAL=RGBColor(0x00,0xC9,0xA7); SLATE=RGBColor(0x64,0x6E,0x83)
CHARCOAL=RGBColor(0x1E,0x1E,0x2E); WHITE=RGBColor(0xFF,0xFF,0xFF)
HEADER_BG="0F172A"; ROW_ALT="F1F5F9"; ROW_WHITE="FFFFFF"
BORDER="CBD5E1"; GREEN_BG="ECFDF5"; AMBER_BG="FFFBEB"; RED_BG="FEF2F2"

# ── EXCHANGE RATE ─────────────────────────────────────────────────────
R = 95.0  # 1 USD = ₹95

# ── API PRICING (June 2026) ──────────────────────────────────────────
# LLM — per 1M tokens (USD, converted to INR)
GEMINI_IN=0.30; GEMINI_OUT=2.50
GROQ_IN=0.59;   GROQ_OUT=0.79

# Voice — Sarvam AI (INR native)
SARVAM_STT_HR  = 30.0    # ₹30/hour
SARVAM_TTS_10K = 15.0    # ₹15/10K chars (Bulbul v2)

# Voice — Google Cloud (USD → INR)
GCLOUD_STT_MIN_USD = 0.016    # $/min
GCLOUD_TTS_1M_USD  = 16.00    # $/1M chars
GCLOUD_STT_MIN = GCLOUD_STT_MIN_USD * R   # ₹/min
GCLOUD_TTS_1M  = GCLOUD_TTS_1M_USD * R    # ₹/1M chars

# ── SCENARIOS ─────────────────────────────────────────────────────────
RQ=5; IT=500; OT=300; SS=15; TC=300; MO=30

A_U=1000; A_V=300; A_T=700; A_RQ=A_U*RQ; A_VR=A_V*RQ
B_U=500;  B_V=150; B_T=350; B_RQ=B_U*RQ; B_VR=B_V*RQ

# ── CALCULATIONS (all in ₹) ──────────────────────────────────────────
def llm(reqs, i_usd, o_usd):
    """LLM cost in ₹"""
    inp = reqs*IT; out = reqs*OT
    return ((inp/1e6)*i_usd + (out/1e6)*o_usd) * R

def sarvam_voice(vreqs):
    """Sarvam STT+TTS cost in ₹ (native INR)"""
    stt = (vreqs*SS/3600) * SARVAM_STT_HR
    tts = (vreqs*TC/10_000) * SARVAM_TTS_10K
    return stt, tts, stt+tts

def gcloud_voice(vreqs):
    """Google Cloud STT+TTS cost in ₹"""
    stt = (vreqs*SS/60) * GCLOUD_STT_MIN
    tts = (vreqs*TC/1e6) * GCLOUD_TTS_1M
    return stt, tts, stt+tts

# Scenario A
a_gem = llm(A_RQ, GEMINI_IN, GEMINI_OUT)
a_groq = llm(A_RQ, GROQ_IN, GROQ_OUT)
a_ss, a_st, a_sv = sarvam_voice(A_VR)
a_gs, a_gt, a_gv = gcloud_voice(A_VR)

# Scenario B
b_gem = llm(B_RQ, GEMINI_IN, GEMINI_OUT)
b_groq = llm(B_RQ, GROQ_IN, GROQ_OUT)
b_ss, b_st, b_sv = sarvam_voice(B_VR)
b_gs, b_gt, b_gv = gcloud_voice(B_VR)

# Stacks (all ₹)
stk = {
    "Gemini + Sarvam":  (a_gem, a_sv, a_gem+a_sv, b_gem, b_sv, b_gem+b_sv),
    "Gemini + Google":  (a_gem, a_gv, a_gem+a_gv, b_gem, b_gv, b_gem+b_gv),
    "Groq + Sarvam":    (a_groq, a_sv, a_groq+a_sv, b_groq, b_sv, b_groq+b_sv),
    "Groq + Google":    (a_groq, a_gv, a_groq+a_gv, b_groq, b_gv, b_groq+b_gv),
}
stk_a = sorted(stk.items(), key=lambda x: x[1][2])   # sort by Scenario A total
stk_b = sorted(stk.items(), key=lambda x: x[1][5])

best_a = stk_a[0]; worst_a = stk_a[-1]
best_b = stk_b[0]; worst_b = stk_b[-1]

# ── FORMATTERS ────────────────────────────────────────────────────────
def r2(v): return f"₹{v:,.2f}"
def r0(v): return f"₹{v:,.0f}"
def d2(v): return f"${v:,.2f}"
def d4(v): return f"${v:,.4f}"
def both(v): return f"₹{v:,.2f}  (${v/R:,.2f})"
def both4(v): return f"₹{v:,.2f}  (${v/R:,.4f})"

# ══════════════════════════════════════════════════════════════════════
# DOCUMENT
# ══════════════════════════════════════════════════════════════════════
doc = Document()
sec = doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=Cm(2.0); sec.bottom_margin=Cm(2.0)
sec.left_margin=Cm(2.5); sec.right_margin=Cm(2.5)

st=doc.styles['Normal']
st.font.name='Calibri'; st.font.size=Pt(10.5); st.font.color.rgb=SLATE
st.paragraph_format.space_after=Pt(6); st.paragraph_format.line_spacing=1.15

for lv,(sz,col,sp) in {1:(26,NAVY,20),2:(17,BLUE,14),3:(13,CHARCOAL,10)}.items():
    h=doc.styles[f'Heading {lv}']
    h.font.name='Calibri'; h.font.size=Pt(sz); h.font.color.rgb=col; h.font.bold=True
    h.paragraph_format.space_before=Pt(sp); h.paragraph_format.space_after=Pt(6)

# ── HELPERS ───────────────────────────────────────────────────────────
def rule(d):
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(6)
    p._element.get_or_add_pPr().append(parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="1" w:color="{BORDER}"/></w:pBdr>'))

def shade(c,clr):
    c._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{clr}" w:val="clear"/>'))

def bord(c, clr=BORDER, s="4"):
    c._tc.get_or_add_tcPr().append(parse_xml(
        f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="{s}" w:space="0" w:color="{clr}"/>'
        f'<w:left w:val="single" w:sz="{s}" w:space="0" w:color="{clr}"/>'
        f'<w:bottom w:val="single" w:sz="{s}" w:space="0" w:color="{clr}"/>'
        f'<w:right w:val="single" w:sz="{s}" w:space="0" w:color="{clr}"/></w:tcBorders>'))

def fmt(table, hdrs, rows, widths=None):
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    tbl=table._tbl
    tblPr=tbl.tblPr if tbl.tblPr is not None else parse_xml(f'<w:tblPr {nsdecls("w")}/>')
    tblPr.append(parse_xml(f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'<w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'<w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'<w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'<w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'<w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/></w:tblBorders>'))
    for i,c in enumerate(table.rows[0].cells):
        shade(c,HEADER_BG); bord(c,HEADER_BG,"6")
        p=c.paragraphs[0]; p.text=""
        rn=p.add_run(hdrs[i]); rn.bold=True; rn.font.size=Pt(9.5); rn.font.color.rgb=WHITE; rn.font.name='Calibri'
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(5)
    for ri,rd in enumerate(rows):
        bg=ROW_ALT if ri%2==0 else ROW_WHITE
        for ci,c in enumerate(table.rows[ri+1].cells):
            shade(c,bg); bord(c,BORDER)
            p=c.paragraphs[0]; p.text=""
            rn=p.add_run(str(rd[ci])); rn.font.size=Pt(9.5); rn.font.color.rgb=CHARCOAL; rn.font.name='Calibri'
            p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(3)
    if widths:
        for row in table.rows:
            for i,c in enumerate(row.cells): c.width=Inches(widths[i])

def bul(d,txt,pfx=None):
    p=d.add_paragraph(style='List Bullet')
    if pfx:
        rb=p.add_run(pfx); rb.bold=True; rb.font.name='Calibri'; rb.font.size=Pt(10.5); rb.font.color.rgb=CHARCOAL
        rn=p.add_run(txt); rn.font.name='Calibri'; rn.font.size=Pt(10.5); rn.font.color.rgb=SLATE
    else:
        rn=p.add_run(txt); rn.font.name='Calibri'; rn.font.size=Pt(10.5); rn.font.color.rgb=SLATE

def body(d,txt):
    p=d.add_paragraph(txt)
    for rn in p.runs: rn.font.name='Calibri'; rn.font.size=Pt(10.5); rn.font.color.rgb=SLATE
    return p

def callout(d,title,value,bg=ROW_ALT):
    p=d.add_paragraph()
    p.paragraph_format.space_before=Pt(6); p.paragraph_format.space_after=Pt(6)
    p._element.get_or_add_pPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg}" w:val="clear"/>'))
    rt=p.add_run(f"  {title}: "); rt.bold=True; rt.font.name='Calibri'; rt.font.size=Pt(11); rt.font.color.rgb=NAVY
    rv=p.add_run(value); rv.bold=True; rv.font.name='Calibri'; rv.font.size=Pt(13); rv.font.color.rgb=BLUE

# ═══════════════════════════════════════════════════════════════════════
#  COVER
# ═══════════════════════════════════════════════════════════════════════
for _ in range(4): doc.add_paragraph()
tp=doc.add_paragraph(); tp.alignment=WD_ALIGN_PARAGRAPH.CENTER
tr=tp.add_run("API Costing\nQuotation"); tr.bold=True; tr.font.size=Pt(36); tr.font.color.rgb=NAVY; tr.font.name='Calibri'
doc.add_paragraph()
sp=doc.add_paragraph(); sp.alignment=WD_ALIGN_PARAGRAPH.CENTER
sr=sp.add_run("AI Application — Daily & Monthly Infrastructure Cost Projection"); sr.font.size=Pt(13); sr.font.color.rgb=BLUE; sr.font.name='Calibri'
doc.add_paragraph()
mp=doc.add_paragraph(); mp.alignment=WD_ALIGN_PARAGRAPH.CENTER
mr=mp.add_run(f"{datetime.date.today().strftime('%B %d, %Y')}  ·  v4.1  ·  CKP Engineering  ·  Confidential"); mr.font.size=Pt(10); mr.font.color.rgb=SLATE; mr.font.name='Calibri'
rp=doc.add_paragraph(); rp.alignment=WD_ALIGN_PARAGRAPH.CENTER
rr=rp.add_run(f"All amounts in Indian Rupees (₹)  ·  1 USD = ₹{R:.0f}"); rr.font.size=Pt(10); rr.font.color.rgb=TEAL; rr.bold=True; rr.font.name='Calibri'

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════
#  TOC
# ═══════════════════════════════════════════════════════════════════════
doc.add_heading("Contents", level=1); rule(doc)
for item in [
    "1   Executive Summary",
    "2   Usage Architecture Breakdown",
    "3   LLM Cost Analysis — Gemini Flash vs. Groq",
    "4   Voice Cost Analysis — Sarvam AI vs. Google Cloud",
    "5   Total Projected Infrastructure Cost",
    "6   Architectural Recommendation",
    "7   Development Quotation & Delivery Plan",
    "     Appendix",
]:
    p=doc.add_paragraph(item); p.paragraph_format.space_after=Pt(3)
    for rn in p.runs: rn.font.size=Pt(11); rn.font.color.rgb=NAVY; rn.font.name='Calibri'
doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════
#  1. EXECUTIVE SUMMARY
# ═══════════════════════════════════════════════════════════════════════
doc.add_heading("1   Executive Summary", level=1); rule(doc)
body(doc,
    "This document evaluates API infrastructure costs for a conversational AI application "
    "with voice capabilities. Two daily load scenarios are modelled — a maximum peak load "
    "(1,000 users) and an average daily load (500 users).")
body(doc, "Two architectural decisions are evaluated:")
bul(doc, " Google Gemini 2.5 Flash vs. Groq Llama 3.3 70B", "LLM Engine:")
bul(doc, " Sarvam AI (STT + Bulbul v2 TTS) vs. Google Cloud (STT V2 + Neural2 TTS)", "Voice Engine:")
body(doc, f"All costs are presented in Indian Rupees (₹). USD-priced services are converted at 1 USD = ₹{R:.0f}. "
     "Sarvam AI is natively priced in INR — no conversion needed.")
doc.add_paragraph()

# ═══════════════════════════════════════════════════════════════════════
#  2. USAGE ARCHITECTURE
# ═══════════════════════════════════════════════════════════════════════
doc.add_heading("2   Usage Architecture Breakdown", level=1); rule(doc)

doc.add_heading("2.1   Scenario Definitions", level=2)
t=doc.add_table(rows=7, cols=3)
fmt(t, ["Parameter","Scenario A — Max","Scenario B — Average"], [
    ["Daily Users",       f"{A_U:,}",          f"{B_U:,}"],
    ["Voice Users",       f"{A_V:,} (30%)",    f"{B_V:,} (30%)"],
    ["Text-Only Users",   f"{A_T:,} (70%)",    f"{B_T:,} (70%)"],
    ["Requests/User",     "5",                 "5"],
    ["Total API Requests",f"{A_RQ:,}",         f"{B_RQ:,}"],
    ["Voice Requests",    f"{A_VR:,}",         f"{B_VR:,}"],
], [2.2,2.2,2.2])

doc.add_heading("2.2   Per-Request Metrics", level=2)
t2=doc.add_table(rows=5, cols=3)
fmt(t2, ["Metric","Value","Applies To"], [
    ["Input Tokens",   f"{IT:,}",          "All requests"],
    ["Output Tokens",  f"{OT:,}",          "All requests"],
    ["STT Audio",      f"{SS} seconds",    "Voice requests only"],
    ["TTS Characters", f"{TC} characters", "Voice requests only"],
], [1.8,1.8,3.0])

doc.add_heading("2.3   Aggregate Volumes", level=2)
ai=A_RQ*IT; ao=A_RQ*OT; bi=B_RQ*IT; bo=B_RQ*OT
asc=A_VR*SS; bsc=B_VR*SS; atc=A_VR*TC; btc=B_VR*TC
t3=doc.add_table(rows=5, cols=3)
fmt(t3, ["Volume","Scenario A","Scenario B"], [
    ["Input Tokens",   f"{ai:,} (2.5M)",     f"{bi:,} (1.25M)"],
    ["Output Tokens",  f"{ao:,} (1.5M)",     f"{bo:,} (0.75M)"],
    ["STT Audio",      f"{asc:,} sec ({asc/3600:.2f} hrs)", f"{bsc:,} sec ({bsc/3600:.2f} hrs)"],
    ["TTS Characters", f"{atc:,}",           f"{btc:,}"],
], [2.2,2.2,2.2])

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════
#  3. LLM COST
# ═══════════════════════════════════════════════════════════════════════
doc.add_heading("3   LLM Cost Analysis — Gemini Flash vs. Groq", level=1); rule(doc)
body(doc, "Comparing Google Gemini 2.5 Flash against Groq Llama 3.3 70B. "
     "Both are priced in USD and converted to ₹.")

doc.add_heading("3.1   Unit Pricing", level=2)
t4=doc.add_table(rows=3, cols=5)
fmt(t4, ["Provider","Input ($/1M)","Output ($/1M)","Input (₹/1M)","Output (₹/1M)"], [
    ["Gemini 2.5 Flash", f"${GEMINI_IN:.2f}", f"${GEMINI_OUT:.2f}", r2(GEMINI_IN*R), r2(GEMINI_OUT*R)],
    ["Groq Llama 3.3 70B", f"${GROQ_IN:.2f}", f"${GROQ_OUT:.2f}", r2(GROQ_IN*R), r2(GROQ_OUT*R)],
], [2.0,1.1,1.1,1.2,1.2])

doc.add_heading("3.2   Daily & Monthly Cost Comparison (₹)", level=2)
sav_a = f"{((a_gem-a_groq)/a_gem)*100:.0f}%" if a_groq<a_gem else "—"
sav_b = f"{((b_gem-b_groq)/b_gem)*100:.0f}%" if b_groq<b_gem else "—"

t5=doc.add_table(rows=5, cols=4)
fmt(t5, ["","Gemini Flash (₹)","Groq 70B (₹)","Groq Savings"], [
    ["A — Daily (5,000 reqs)",  r2(a_gem),         r2(a_groq),         sav_a],
    ["A — Monthly (30 days)",   r2(a_gem*MO),      r2(a_groq*MO),      sav_a],
    ["B — Daily (2,500 reqs)",  r2(b_gem),         r2(b_groq),         sav_b],
    ["B — Monthly (30 days)",   r2(b_gem*MO),      r2(b_groq*MO),      sav_b],
], [2.4,1.6,1.6,1.0])

winner_llm = "Groq 70B" if a_groq<a_gem else "Gemini Flash"
wl = min(a_groq,a_gem)
callout(doc, "🏆 Lowest LLM Cost", f"{winner_llm} — {both4(wl)}/day (Max)", GREEN_BG)

doc.add_heading("3.3   Observations", level=2)
if a_groq<a_gem:
    bul(doc, f" Groq 70B is {a_gem/a_groq:.1f}× cheaper: {r2(a_groq)} vs {r2(a_gem)}/day.", "Cost:")
    bul(doc, f" Monthly saving at max load: {r2((a_gem-a_groq)*MO)}.", "Savings:")
bul(doc, " Gemini Flash — best multimodal reasoning & Google ecosystem.", "Quality:")
bul(doc, " Groq LPU™ — <100ms/token latency, critical for real-time chat.", "Latency:")

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════
#  4. VOICE COST
# ═══════════════════════════════════════════════════════════════════════
doc.add_heading("4   Voice Cost Analysis — Sarvam AI vs. Google Cloud", level=1); rule(doc)
body(doc, "Comparing Sarvam AI (STT + Bulbul v2 TTS, native ₹ pricing) against "
     "Google Cloud (STT V2 + Neural2 TTS, USD → ₹ converted).")

doc.add_heading("4.1   Unit Pricing", level=2)
t6=doc.add_table(rows=3, cols=4)
fmt(t6, ["Service","Sarvam AI (₹)","Google Cloud ($)","Google Cloud (₹)"], [
    ["STT", f"₹{SARVAM_STT_HR:.0f}/hour", f"${GCLOUD_STT_MIN_USD}/min", f"{r2(GCLOUD_STT_MIN)}/min"],
    ["TTS", f"₹{SARVAM_TTS_10K:.0f}/10K chars", f"${GCLOUD_TTS_1M_USD:.2f}/1M chars", f"{r2(GCLOUD_TTS_1M)}/1M chars"],
], [1.2,1.8,1.6,1.8])

doc.add_heading("4.2   Scenario A — Max Load (1,500 voice reqs)", level=2)
t7a=doc.add_table(rows=5, cols=4)
fmt(t7a, ["Component","Sarvam AI (₹)","Google Cloud (₹)","Cheaper"], [
    ["STT (daily)",      r2(a_ss), r2(a_gs), "Sarvam" if a_ss<a_gs else "Google"],
    ["TTS (daily)",      r2(a_st), r2(a_gt), "Sarvam" if a_st<a_gt else "Google"],
    ["TOTAL (daily)",    r2(a_sv), r2(a_gv), "Sarvam" if a_sv<a_gv else "Google"],
    ["TOTAL (monthly)",  r2(a_sv*MO), r2(a_gv*MO), "Sarvam" if a_sv<a_gv else "Google"],
], [1.8,1.6,1.6,1.2])

doc.add_heading("4.3   Scenario B — Average Load (750 voice reqs)", level=2)
t7b=doc.add_table(rows=5, cols=4)
fmt(t7b, ["Component","Sarvam AI (₹)","Google Cloud (₹)","Cheaper"], [
    ["STT (daily)",      r2(b_ss), r2(b_gs), "Sarvam" if b_ss<b_gs else "Google"],
    ["TTS (daily)",      r2(b_st), r2(b_gt), "Sarvam" if b_st<b_gt else "Google"],
    ["TOTAL (daily)",    r2(b_sv), r2(b_gv), "Sarvam" if b_sv<b_gv else "Google"],
    ["TOTAL (monthly)",  r2(b_sv*MO), r2(b_gv*MO), "Sarvam" if b_sv<b_gv else "Google"],
], [1.8,1.6,1.6,1.2])

wv = "Sarvam AI" if a_sv<a_gv else "Google Cloud"
wvc = min(a_sv,a_gv)
callout(doc, "🏆 Lowest Voice Cost", f"{wv} — {both4(wvc)}/day (Max)", GREEN_BG)

doc.add_heading("4.4   Observations", level=2)
if a_ss<a_gs: bul(doc, f" Sarvam STT: {r2(a_ss)} vs Google {r2(a_gs)}/day.", "STT:")
else:         bul(doc, f" Google STT: {r2(a_gs)} vs Sarvam {r2(a_ss)}/day.", "STT:")
if a_st<a_gt: bul(doc, f" Sarvam TTS: {r2(a_st)} vs Google {r2(a_gt)}/day.", "TTS:")
else:         bul(doc, f" Google TTS: {r2(a_gt)} vs Sarvam {r2(a_st)}/day.", "TTS:")
bul(doc, " Sarvam Bulbul v2 — best-in-class Indic language naturalness.", "Language:")
bul(doc, " Google Cloud — 40+ languages, enterprise SLAs.", "Global:")

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════
#  5. TOTAL PROJECTED COST
# ═══════════════════════════════════════════════════════════════════════
doc.add_heading("5   Total Projected Infrastructure Cost", level=1); rule(doc)
body(doc, "All four stack combinations (2 LLMs × 2 Voice engines), "
     "daily and monthly costs in ₹ for both scenarios.")

doc.add_heading("5.1   Scenario A — Max Load (1,000 users/day)", level=2)
t8a=doc.add_table(rows=len(stk_a)+1, cols=5)
r8a=[]
for nm,(al,av,at,_,_,_) in stk_a:
    r8a.append([nm, r2(al), r2(av), r2(at), r2(at*MO)])
fmt(t8a, ["Stack","LLM/Day (₹)","Voice/Day (₹)","Total/Day (₹)","Total/Month (₹)"], r8a, [2.0,1.2,1.2,1.4,1.4])

doc.add_heading("5.2   Scenario B — Average Load (500 users/day)", level=2)
t8b=doc.add_table(rows=len(stk_b)+1, cols=5)
r8b=[]
for nm,(_,_,_,bl,bv,bt) in stk_b:
    r8b.append([nm, r2(bl), r2(bv), r2(bt), r2(bt*MO)])
fmt(t8b, ["Stack","LLM/Day (₹)","Voice/Day (₹)","Total/Day (₹)","Total/Month (₹)"], r8b, [2.0,1.2,1.2,1.4,1.4])

doc.add_paragraph()
callout(doc, "💚 Best Value (Max)", f"{best_a[0]} — {r2(best_a[1][2])}/day · {r2(best_a[1][2]*MO)}/month", GREEN_BG)
callout(doc, "💚 Best Value (Avg)", f"{best_b[0]} — {r2(best_b[1][5])}/day · {r2(best_b[1][5]*MO)}/month", GREEN_BG)
callout(doc, "🔶 Max Cost (Max)", f"{worst_a[0]} — {r2(worst_a[1][2])}/day · {r2(worst_a[1][2]*MO)}/month", AMBER_BG)

doc.add_heading("5.3   Monthly Cost Range", level=2)
t9=doc.add_table(rows=3, cols=4)
fmt(t9, ["","Best (₹/month)","Worst (₹/month)","Spread (₹)"], [
    ["Scenario A (Max)", f"{best_a[0]}: {r2(best_a[1][2]*MO)}", f"{worst_a[0]}: {r2(worst_a[1][2]*MO)}", r2((worst_a[1][2]-best_a[1][2])*MO)],
    ["Scenario B (Avg)", f"{best_b[0]}: {r2(best_b[1][5]*MO)}", f"{worst_b[0]}: {r2(worst_b[1][5]*MO)}", r2((worst_b[1][5]-best_b[1][5])*MO)],
], [1.6,2.0,2.0,1.2])

doc.add_heading("5.4   USD Equivalents (Reference)", level=2)
body(doc, f"For international reference (1 USD = ₹{R:.0f}):")

t10=doc.add_table(rows=len(stk_a)+1, cols=4)
r10=[]
for nm,(al,av,at,bl,bv,bt) in stk_a:
    r10.append([nm, d2(at/R), d2(at*MO/R), d2(bt*MO/R)])
fmt(t10, ["Stack","A: Daily ($)","A: Monthly ($)","B: Monthly ($)"], r10, [2.0,1.6,1.6,1.6])

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════
#  6. RECOMMENDATION
# ═══════════════════════════════════════════════════════════════════════
doc.add_heading("6   Architectural Recommendation", level=1); rule(doc)

best_nm = best_a[0]
best_vals = best_a[1]
callout(doc, "✅  Recommended Stack", best_nm, GREEN_BG)

doc.add_heading("6.1   Recommended Stack Breakdown", level=2)
parts = best_nm.split(" + ")
llm_nm = parts[0]; voice_nm = parts[1]
llm_d = best_vals[0]; voice_d = best_vals[1]; total_d = best_vals[2]

t_rec=doc.add_table(rows=4, cols=4)
fmt(t_rec, ["Component","Provider","Daily (₹)","Monthly (₹)"], [
    ["LLM",    llm_nm,   r2(llm_d),   r2(llm_d*MO)],
    ["Voice",  voice_nm, r2(voice_d), r2(voice_d*MO)],
    ["TOTAL",  best_nm,  r2(total_d), r2(total_d*MO)],
], [1.4,1.8,1.6,1.6])

doc.add_heading("6.2   Why This Stack", level=2)
if "Groq" in best_nm:
    bul(doc, f" Groq 70B at {r2(a_groq)}/day — strong 70B reasoning with <100ms/token latency.", "Speed + Value:")
if "Sarvam" in best_nm:
    bul(doc, f" Sarvam AI at {r2(a_sv)}/day — cheapest voice option with best Indic language quality.", "Voice Leader:")
if "Google" in best_nm:
    bul(doc, f" Google Cloud at {r2(a_gv)}/day — enterprise SLAs, 40+ languages.", "Enterprise:")
if "Gemini" in best_nm:
    bul(doc, f" Gemini Flash at {r2(a_gem)}/day — Google's best multimodal reasoning.", "Quality:")

doc.add_heading("6.3   Strategic Guidance", level=2)
bul(doc, f" Daily cost range: {r2(best_a[1][2])} to {r2(worst_a[1][2])} at max load.", "Cost Envelope:")
bul(doc, f" Monthly range: {r2(best_a[1][2]*MO)} to {r2(worst_a[1][2]*MO)} — all stacks are very affordable.", "Monthly:")
bul(doc, " Route simple queries to a lighter model (Groq 8B at ₹4.75/₹7.60 per 1M tokens) for further savings.", "Smart Routing:")
bul(doc, " Both Groq and Gemini offer ~50% batch API discounts for async workloads.", "Batch Pricing:")
bul(doc, " For Indic-language apps, Sarvam AI is the clear voice winner in both cost and quality.", "Language:")
bul(doc, f" At max load, monthly ceiling is only {r2(worst_a[1][2]*MO)} ({d2(worst_a[1][2]*MO/R)}). "
     "API costs are not a bottleneck.", "Bottom Line:")

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════
#  7. DEVELOPMENT QUOTATION & DELIVERY PLAN
# ═══════════════════════════════════════════════════════════════════════
doc.add_heading("7   Development Quotation & Delivery Plan", level=1); rule(doc)

body(doc, "We propose a fixed-price development timeline of 3 months, aligning with the college website's progress. "
     "This covers the deployment of the chatbot across three departments (Engineering, Commerce, and Pharmacy) "
     "and its integration into the main portal. The total quote is broken down below:")

t_dev = doc.add_table(rows=5, cols=3)
fmt(t_dev, ["Milestone", "Scope / Deliverable Modules", "Cost (₹)"], [
    ["Milestone 1: Core System & Engineering Bot", "Backend RAG pipeline, local vectors, conversation memory, and React chat interface UI.", r0(45000)],
    ["Milestone 2: Commerce Department Setup", "Curation of localized dataset, separate vector indices, and custom department configurations.", r0(15000)],
    ["Milestone 3: Pharmacy Department Setup", "Curation of localized dataset, separate vector indices, and custom department configurations.", r0(15000)],
    ["Milestone 4: Cloud Hosting & Integration", "AWS deployment, SSL domain mapping, and embedding iframe integration on the main portal.", r0(15000)],
], [2.0, 3.5, 1.5])

total_dev = 45000 + 15000 + 15000 + 15000
callout(doc, "💎 Total Development Cost", f"{r0(total_dev)} (Fixed-Price)", ROW_ALT)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════
#  APPENDIX
# ═══════════════════════════════════════════════════════════════════════
doc.add_heading("Appendix — Pricing Sources", level=1); rule(doc)
body(doc, "Pricing from official vendor docs, June 2026.")
bul(doc, " https://ai.google.dev/pricing", "Gemini:")
bul(doc, " https://groq.com/pricing", "Groq:")
bul(doc, " https://www.sarvam.ai/pricing", "Sarvam AI:")
bul(doc, " https://cloud.google.com/speech-to-text/pricing", "Google STT:")
bul(doc, " https://cloud.google.com/text-to-speech/pricing", "Google TTS:")
doc.add_paragraph()
bul(doc, f" 1 USD = ₹{R:.0f} (June 2026)", "Exchange Rate:")
bul(doc, " Sarvam AI priced natively in ₹ — no conversion applied", "Note:")
doc.add_paragraph()
dp=doc.add_paragraph()
rd=dp.add_run("Disclaimer: Pricing based on public rate cards. Actual costs may vary. For planning only.")
rd.font.size=Pt(9); rd.font.color.rgb=SLATE; rd.font.name='Calibri'; rd.italic=True

# Footer
ft=sec.footer; ft.is_linked_to_previous=False
fp=ft.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
fr=fp.add_run("CKP Engineering  ·  API Costing Quotation v4.1  ·  All amounts in ₹  ·  Confidential")
fr.font.size=Pt(8); fr.font.color.rgb=SLATE; fr.font.name='Calibri'

# ═══════════════════════════════════════════════════════════════════════
#  SAVE
# ═══════════════════════════════════════════════════════════════════════
out="/Users/sujalvachhani/Desktop/CKP-Engineering-bot-main/API_Costing_Quotation.docx"
doc.save(out)

print(f"\n✅ Saved: {out}")
print(f"\n{'='*70}")
print(f"  API COSTING v4.1 — ALL AMOUNTS IN ₹ (INR)")
print(f"  1 USD = ₹{R:.0f}  |  Sarvam TTS: Bulbul v2 @ ₹15/10K chars")
print(f"{'='*70}")
print(f"\n  SCENARIO A (MAX — {A_U:,} users, {A_RQ:,} reqs/day)")
print(f"  {'─'*55}")
print(f"  LLM:   Gemini Flash    {r2(a_gem):>14}/day   {r2(a_gem*MO):>14}/mo")
print(f"         Groq 70B        {r2(a_groq):>14}/day   {r2(a_groq*MO):>14}/mo")
print(f"  Voice: Sarvam AI       {r2(a_sv):>14}/day   {r2(a_sv*MO):>14}/mo")
print(f"         Google Cloud    {r2(a_gv):>14}/day   {r2(a_gv*MO):>14}/mo")
print(f"\n  SCENARIO B (AVG — {B_U:,} users, {B_RQ:,} reqs/day)")
print(f"  {'─'*55}")
print(f"  LLM:   Gemini Flash    {r2(b_gem):>14}/day   {r2(b_gem*MO):>14}/mo")
print(f"         Groq 70B        {r2(b_groq):>14}/day   {r2(b_groq*MO):>14}/mo")
print(f"  Voice: Sarvam AI       {r2(b_sv):>14}/day   {r2(b_sv*MO):>14}/mo")
print(f"         Google Cloud    {r2(b_gv):>14}/day   {r2(b_gv*MO):>14}/mo")
print(f"\n  FULL STACKS — SCENARIO A (daily → monthly):")
for nm,(al,av,at,bl,bv,bt) in stk_a:
    print(f"    {nm:24s}  {r2(at):>14}/day   {r2(at*MO):>14}/mo")
print(f"\n  ✅ BEST:  {best_a[0]:24s}  {r2(best_a[1][2])}/day  →  {r2(best_a[1][2]*MO)}/mo")
print(f"  🔶 MAX:   {worst_a[0]:24s}  {r2(worst_a[1][2])}/day  →  {r2(worst_a[1][2]*MO)}/mo")
print(f"{'='*70}")
